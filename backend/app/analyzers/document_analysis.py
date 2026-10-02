"""Document Analysis — category weight 10 %.

Extracts text from uploaded PDF or document offer letters using pypdf.
Reviews document structure, payment / deposit clauses, fake stamp annotations,
and non-standard HR onboarding procedures.

An uploaded **certificate image** (PNG/JPG) or a scanned PDF additionally gets a
visual check from the fine-tuned ONNX forgery model
(:mod:`app.services.certificate_forensics`).  That model runs on 287 training
images whose class split is partly explained by file size and orientation, so its
verdict is deliberately capped at ``Severity.MEDIUM``: a high-severity factor
would force the engine's 65-point floor and let this single weak signal flip the
recommendation to *DON'T APPLY*.
"""

from __future__ import annotations

import io
import logging
import re

from app.models.schemas import (
    CategoryResult,
    RiskCategory,
    RiskFactor,
    Severity,
)
from app.services import certificate_forensics

logger = logging.getLogger(__name__)

# Image formats the forgery model can decode.  Binary image bytes must never be
# decoded as text — that produces mojibake which trips the clause regexes below
# and invents risk that the image does not contain.
_IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff")

# Score deltas for the visual forgery signal.  Kept small because the model is a
# secondary signal: at the 0.10 document weight these contribute ±2.2 / −0.8
# points to the final 0–100 score, far too little to move a band on its own.
_FORGERY_FLAGGED_DELTA = 22.0
_FORGERY_CLEAR_DELTA = -8.0


def extract_document_text(document_bytes: bytes | None, filename: str | None = None) -> str:
    """Extract readable text from PDF bytes or raw text document."""
    if not document_bytes:
        return ""

    # An uploaded image is scored by the visual model, not by text extraction.
    if filename and filename.lower().endswith(_IMAGE_SUFFIXES):
        return ""

    text = ""
    # Try pypdf extraction if PDF
    if filename and filename.lower().endswith(".pdf") or document_bytes.startswith(b"%PDF"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(document_bytes))
            for page in reader.pages:
                page_text = page.extract_text() or ""
                text += f"\n{page_text}"
        except Exception as exc:
            logger.warning("pypdf extraction failed on %s: %s", filename, exc)

    # .docx: binary OOXML — need a real parser (zip archive of XML parts).
    if filename and filename.lower().endswith(".docx"):
        try:
            import docx as _docx  # python-docx
            _blob = _docx.Document(io.BytesIO(document_bytes))
            text = "\n".join(
                paragraph.text for paragraph in _blob.paragraphs if paragraph.text.strip()
            )
            # Tables are a common place for fees / clauses in offer letters.
            for table in _blob.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            text += "\n" + cell_text
            if text.strip():
                return text.strip()
        except Exception as exc:
            logger.warning("docx extraction failed on %s: %s", filename, exc)

    # If text is still empty or wasn't a PDF, attempt utf-8 / latin-1 decoding
    if not text.strip():
        try:
            text = document_bytes.decode("utf-8", errors="ignore")
        except Exception:
            text = document_bytes.decode("latin-1", errors="ignore")

    return text.strip()


def extract_pdf_metadata(document_bytes: bytes | None, filename: str | None = None) -> dict[str, str]:
    """Extract forensic metadata (Producer, Creator, Author) from PDF bytes."""
    if not document_bytes:
        return {}
    if (filename and filename.lower().endswith(".pdf")) or document_bytes.startswith(b"%PDF"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(document_bytes))
            meta = reader.metadata or {}
            return {
                "producer": str(meta.get("/Producer", "") or "").strip(),
                "creator": str(meta.get("/Creator", "") or "").strip(),
                "author": str(meta.get("/Author", "") or "").strip(),
            }
        except Exception as exc:
            logger.debug("PDF metadata extraction failed: %s", exc)
    return {}


async def analyze(
    document_bytes: bytes | None = None,
    document_filename: str | None = None,
    document_text: str | None = None,
    **kwargs,
) -> CategoryResult:
    """Analyse uploaded offer letter or PDF document.

    Returns:
        A :class:`CategoryResult` with score (0–100), risk factors, and analyzed=True/False.
    """
    raw_text = document_text or extract_document_text(document_bytes, document_filename)
    pdf_meta = extract_pdf_metadata(document_bytes, document_filename)

    # Visual forgery model (ONNX).  Run it before the "nothing to analyse" check
    # so a lone certificate *image* still produces a result: image uploads yield
    # no extractable text and no PDF metadata by design.
    forgery_signal = certificate_forensics.analyse_document(document_bytes, document_filename)

    if not raw_text.strip() and not pdf_meta and not forgery_signal.available:
        return CategoryResult(score=0.0, risk_factors=[], analyzed=False)


    text_lower = raw_text.lower()
    risk_factors: list[RiskFactor] = []
    base_score = 0.0

    # 1. Deposit / Hardware Purchase clause in Offer Letter
    deposit_pattern = re.search(
        r"(?:refundable|security|equipment|laptop|training|kit|\s)*(?:deposit|charge|fee|amount)\s*(?:of|is)?\s*(?:rs\.?|inr|₹|\$)?\s*(\d+[\d,]*)",
        text_lower,
    )
    if deposit_pattern:
        base_score += 65.0
        risk_factors.append(
            RiskFactor(
                category=RiskCategory.DOCUMENT_ANALYSIS,
                severity=Severity.HIGH,
                description="Security deposit or equipment purchase clause in offer letter.",
                evidence=(
                    f"Offer letter specifies financial condition: '{deposit_pattern.group(0)}'. "
                    "Legitimate corporate offer letters do not require candidates to pay fees or purchase gear."
                ),
                source="pdf_clause_extractor",
                confidence=0.96,
            )
        )

    # 2. Fake notary / Government stamp claims in private employment contract
    stamp_pattern = re.search(
        r"(?:notarized stamp|govt approved bond|legal agreement bond|ministry of corporate affairs stamp|registered seal)",
        text_lower,
    )
    if stamp_pattern:
        base_score += 35.0
        risk_factors.append(
            RiskFactor(
                category=RiskCategory.DOCUMENT_ANALYSIS,
                severity=Severity.MEDIUM,
                description="Unusual legal stamp or court bond phrasing in offer document.",
                evidence=f"Text references '{stamp_pattern.group(0)}', commonly added to counterfeit offer letters to intimidate freshers.",
                source="counterfeit_document_pattern_checker",
                confidence=0.88,
            )
        )

    # 3. Informal communication channel in official appointment letter
    chat_pattern = re.search(r"(?:telegram|whatsapp|wa\.me/|@gmail\.com|@yahoo\.)", text_lower)
    if chat_pattern:
        base_score += 35.0
        risk_factors.append(
            RiskFactor(
                category=RiskCategory.DOCUMENT_ANALYSIS,
                severity=Severity.MEDIUM,
                description="Offer document directs candidate to informal chat channel or personal email.",
                evidence=f"Document references contact '{chat_pattern.group(0)}' for onboarding and reporting.",
                source="document_channel_inspector",
                confidence=0.90,
            )
        )

    # 4. Forensic PDF metadata inspection (Producer, Creator, Author)
    if pdf_meta:
        creator_combined = f"{pdf_meta.get('creator', '')} {pdf_meta.get('producer', '')} {pdf_meta.get('author', '')}".lower()
        
        # Check consumer graphic design tools
        consumer_tools = ["canva", "photoshop", "coreldraw", "illustrator", "paint.net", "gimp"]
        matched_consumer = next((tool for tool in consumer_tools if tool in creator_combined), None)

        # Check enterprise e-sign & HR platforms
        enterprise_tools = ["docusign", "adobesign", "workday", "bamboohr", "successfactors"]
        matched_enterprise = next((tool for tool in enterprise_tools if tool in creator_combined), None)

        if matched_consumer:
            base_score += 45.0
            tool_name = pdf_meta.get("creator") or pdf_meta.get("producer") or matched_consumer.title()
            risk_factors.append(
                RiskFactor(
                    category=RiskCategory.DOCUMENT_ANALYSIS,
                    severity=Severity.HIGH,
                    description=f"Consumer design tool ({tool_name}) detected in document metadata.",
                    evidence=(
                        f"PDF metadata identifies creation tool as '{tool_name}'. "
                        "Official corporate offer letters are generated via enterprise HR ERPs "
                        "or verified e-sign platforms (Workday, DocuSign), not consumer graphic design tools."
                    ),
                    source="pdf_forensic_metadata_inspector",
                    confidence=0.94,
                )
            )
        elif matched_enterprise:
            base_score = max(base_score - 15.0, 5.0)
            ent_name = pdf_meta.get("producer") or pdf_meta.get("creator") or matched_enterprise.title()
            risk_factors.append(
                RiskFactor(
                    category=RiskCategory.DOCUMENT_ANALYSIS,
                    severity=Severity.LOW,
                    description=f"Verified corporate e-sign / HR platform ({ent_name}) detected.",
                    evidence=f"PDF metadata confirms document was produced using enterprise e-signature infrastructure '{ent_name}'.",
                    source="pdf_forensic_metadata_inspector",
                    confidence=0.92,
                )
            )

    # 5. If document contains legitimate employment terms (provident fund, CTC, gratuity, medical)
    has_legit_clauses = any(
        kw in text_lower for kw in ["provident fund", "pf", "gratuity", "cost to company", "ctc", "leaves", "health insurance"]
    )
    if has_legit_clauses and not deposit_pattern:
        base_score = max(base_score - 10.0, 5.0)
        risk_factors.append(
            RiskFactor(
                category=RiskCategory.DOCUMENT_ANALYSIS,
                severity=Severity.LOW,
                description="Standard formal compensation and benefits structure detected.",
                evidence="Document contains standard corporate employment clauses (PF, benefits, leave policy).",
                source="hr_structure_validator",
                confidence=0.85,
            )
        )

    # 6. Visual forgery model — SECONDARY signal, capped at MEDIUM severity
    #    A HIGH severity factor would trip the risk engine's 65-point floor and
    #    let this model alone force a DON'T APPLY verdict, which its ~0.63
    #    metadata-only baseline and 287-image training set do not justify.
    if forgery_signal.available:
        if forgery_signal.flagged:
            base_score += _FORGERY_FLAGGED_DELTA
            risk_factors.append(
                RiskFactor(
                    category=RiskCategory.DOCUMENT_ANALYSIS,
                    severity=Severity.MEDIUM,
                    description=(
                        "Visual forgery model flags this certificate as possibly altered."
                    ),
                    evidence=(
                        f"Fine-tuned MobileNetV3-Small scored p(forged)="
                        f"{forgery_signal.fake_probability:.2f} against a "
                        f"{forgery_signal.threshold:.2f} threshold. "
                        "This is a secondary signal: the model was trained on a small "
                        "public dataset and must be corroborated by other evidence "
                        "before it affects a decision."
                    ),
                    source="certificate_cnn_onnx",
                    confidence=0.55,
                )
            )
        else:
            base_score = max(base_score + _FORGERY_CLEAR_DELTA, 0.0)
            risk_factors.append(
                RiskFactor(
                    category=RiskCategory.DOCUMENT_ANALYSIS,
                    severity=Severity.LOW,
                    description="Certificates visual layout appears consistent with a genuine document.",
                    evidence=(
                        f"Fine-tuned MobileNetV3-Small scored p(forged)="
                        f"{forgery_signal.fake_probability:.2f}, below the "
                        f"{forgery_signal.threshold:.2f} threshold. "
                        "The absence of a forgery signal is not proof of authenticity."
                    ),
                    source="certificate_cnn_onnx",
                    confidence=0.50,
                )
            )

    final_score = round(min(max(base_score, 0.0), 100.0), 2)
    return CategoryResult(score=final_score, risk_factors=risk_factors, analyzed=True)

