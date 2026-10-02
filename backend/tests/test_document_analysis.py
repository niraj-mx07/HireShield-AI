"""Tests for document analysis module."""

from __future__ import annotations

import pytest
from app.analyzers import document_analysis
from app.models.schemas import (
    CategoryResult,
    Recommendation,
    RiskBand,
    RiskCategory,
    Severity,
)
from app.services.certificate_forensics import ForgerySignal
from app.services.risk_engine import score_assessment


@pytest.mark.asyncio
async def test_scam_text_in_document():
    """Test extracted text from offer letter containing fee demands yields high risk."""
    doc_text = "OFFER LETTER\nTo claim your joining kit, candidate must deposit fee of Rs 5000 refundable security deposit to bank account below."
    result = await document_analysis.analyze(document_text=doc_text)
    assert result.analyzed is True
    assert result.score >= 50.0
    assert any(rf.severity == Severity.HIGH for rf in result.risk_factors)


@pytest.mark.asyncio
async def test_legitimate_document_text():
    """Test legitimate document text returns low risk."""
    doc_text = "OFFER OF EMPLOYMENT\nWe are pleased to offer you the role of Software Engineer at Acme Corp. Annual compensation is $120,000."
    result = await document_analysis.analyze(document_text=doc_text)
    assert result.analyzed is True
    assert result.score < 30.0


@pytest.mark.asyncio
async def test_empty_document_unanalyzed():
    """Test empty document text returns analyzed=False."""
    result = await document_analysis.analyze(document_text="")
    assert result.analyzed is False


@pytest.mark.asyncio
async def test_canva_pdf_metadata_flagged(monkeypatch):
    """Test PDF created in Canva or consumer design tool triggers high severity."""
    monkeypatch.setattr(
        document_analysis,
        "extract_pdf_metadata",
        lambda *args, **kwargs: {"creator": "Canva", "producer": "Canva PDF Exporter", "author": "Anonymous"},
    )
    result = await document_analysis.analyze(
        document_bytes=b"%PDF-mock",
        document_filename="offer_letter.pdf",
        document_text="Offer of employment at Tech Global Solutions",
    )
    assert result.analyzed is True
    assert result.score >= 40.0
    assert any("canva" in rf.description.lower() for rf in result.risk_factors)


# ---------------------------------------------------------------------------
# .docx text extraction (new: OOXML parsing via python-docx)
# ---------------------------------------------------------------------------

def _docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Build a minimal valid .docx byte blob from paragraphs + a table."""
    import io

    import docx

    blob = docx.Document()
    for paragraph in paragraphs:
        blob.add_paragraph(paragraph)
    if table_rows:
        table = blob.add_table(rows=len(table_rows), cols=max(len(r) for r in table_rows))
        for row_idx, row in enumerate(table_rows):
            for col_idx, cell_text in enumerate(row):
                table.cell(row_idx, col_idx).text = cell_text
    buffer = io.BytesIO()
    blob.save(buffer)
    return buffer.getvalue()


def test_docx_paragraphs_extracted():
    """Paragraph text must be extracted from a .docx upload."""
    payload = _docx_bytes(
        paragraphs=[
            "OFFER OF EMPLOYMENT",
            "We are pleased to offer you the role of Software Engineer at Acme Corp.",
            "Annual compensation is $120,000.",
        ]
    )
    extracted = document_analysis.extract_document_text(payload, "offer.docx")
    assert "Software Engineer" in extracted
    assert "Acme Corp" in extracted


def test_docx_table_extracted():
    """Table cells (a common home for fee / clause text) must be extracted too."""
    payload = _docx_bytes(
        paragraphs=[],
        table_rows=[
            ["Field", "Value"],
            ["Company", "Acme Corp"],
            ["CTC", "Rs 15,00,000"],
        ],
    )
    extracted = document_analysis.extract_document_text(payload, "offer.docx")
    assert "Acme Corp" in extracted
    assert "15,00,000" in extracted


def test_docx_non_ascii_safe():
    """Non-ASCII text in the document must survive the decode path."""
    payload = _docx_bytes(paragraphs=["Offer letter for राजा अनुयायी and ₹5,000 deposit."])
    extracted = document_analysis.extract_document_text(payload, "offer.docx")
    assert "राजा अनुयायी" in extracted
    assert "₹5,000" in extracted


def test_docx_raises_gracefully():
    """A corrupt or malformed .docx must not crash — falls back to the generic decode."""
    payload = b"this is not a real docx file"
    extracted = document_analysis.extract_document_text(payload, "offer.docx")
    assert "not a real docx file" in extracted


# ---------------------------------------------------------------------------
# Visual certificate-forgery model (secondary signal)
# ---------------------------------------------------------------------------

def _signal(flagged: bool, probability: float = 0.9, available: bool = True):
    """Build a :class:`ForgerySignal` without running ONNX."""
    return ForgerySignal(
        available=available,
        fake_probability=probability if available else None,
        threshold=0.5,
        flagged=flagged,
        notes=["visual check ran"],
    )


def _png_bytes(size=(96, 72)) -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, (200, 210, 220)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_image_upload_is_not_decoded_as_text():
    """Binary image bytes must never be decoded as text (mojibake trips regexes)."""
    payload = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
    assert document_analysis.extract_document_text(payload, "certificate.png") == ""
    assert document_analysis.extract_document_text(payload, "scan.JPG") == ""


@pytest.mark.asyncio
async def test_flagged_certificate_image_adds_medium_factor(monkeypatch):
    """A flagged image upload must still be analyzed and stay at MEDIUM severity."""
    monkeypatch.setattr(
        document_analysis.certificate_forensics, "analyse_document",
        lambda *a, **k: _signal(flagged=True, probability=0.83),
    )

    result = await document_analysis.analyze(
        document_bytes=_png_bytes(), document_filename="certificate.png"
    )

    assert result.analyzed is True
    assert result.score > 0.0
    cnn = [rf for rf in result.risk_factors if rf.source == "certificate_cnn_onnx"]
    assert len(cnn) == 1
    assert cnn[0].severity == Severity.MEDIUM
    assert "0.83" in cnn[0].evidence


@pytest.mark.asyncio
async def test_clear_certificate_image_adds_low_factor(monkeypatch):
    monkeypatch.setattr(
        document_analysis.certificate_forensics, "analyse_document",
        lambda *a, **k: _signal(flagged=False, probability=0.04),
    )

    result = await document_analysis.analyze(
        document_bytes=_png_bytes(), document_filename="certificate.png"
    )

    assert result.analyzed is True
    cnn = [rf for rf in result.risk_factors if rf.source == "certificate_cnn_onnx"]
    assert len(cnn) == 1
    assert cnn[0].severity == Severity.LOW
    assert result.score == 0.0


@pytest.mark.asyncio
async def test_image_without_deployed_model_is_unanalyzed(monkeypatch):
    """No model + no text + no PDF metadata must return analyzed=False, not a crash."""
    monkeypatch.setattr(
        document_analysis.certificate_forensics, "analyse_document",
        lambda *a, **k: _signal(flagged=False, available=False),
    )

    result = await document_analysis.analyze(
        document_bytes=_png_bytes(), document_filename="certificate.png"
    )

    assert result.analyzed is False
    assert result.risk_factors == []


@pytest.mark.asyncio
async def test_forgery_flag_alone_cannot_force_dont_apply(monkeypatch):
    """Security property: the weak visual model must never force a DON'T APPLY verdict.

    The risk engine forces ``DONT_APPLY`` whenever any factor is HIGH severity and
    raises the score to 65+.  A model trained on 287 images with a 0.63
    metadata-only baseline is nowhere near strong enough to justify that, so the
    CNN factor is capped at MEDIUM — this test locks that in.
    """
    monkeypatch.setattr(
        document_analysis.certificate_forensics, "analyse_document",
        lambda *a, **k: _signal(flagged=True, probability=0.99),
    )

    document_result = await document_analysis.analyze(
        document_bytes=_png_bytes(), document_filename="certificate.png"
    )

    # 1. The factor itself must never claim HIGH severity.
    assert all(rf.severity != Severity.HIGH for rf in document_result.risk_factors)

    # 2. Even as the only signal in the whole assessment, it must not flip the verdict.
    results = {
        category: (
            document_result
            if category == RiskCategory.DOCUMENT_ANALYSIS
            else CategoryResult(score=0.0, risk_factors=[], analyzed=True)
        )
        for category in RiskCategory
    }
    score, band, recommendation, _confidence, _scores, _factors = score_assessment(results)

    assert score < 65.0
    assert band not in (RiskBand.HIGH, RiskBand.VERY_HIGH)
    assert recommendation != Recommendation.DONT_APPLY

